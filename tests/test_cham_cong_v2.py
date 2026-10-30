"""Test nghiem thu cham cong v2 (T1-T25) va phep / giai trinh (P1-P7) - docs/LOGIC_CHAM_CONG.md.
Ca mau: 07:00-17:00, nghi 11:30-12:30, an han 5, cho vao som 60, chong quet trung 10 phut, missing_checkout_hours 4.
CSDL SQLite trong bo nho (T22: file SQLite tam) hoac TEST_DATABASE_URL (MariaDB / PostgreSQL), xem app_helpers.py.
Chay: .venv\\Scripts\\python.exe -m unittest tests.test_cham_cong_v2 -v   (tu thu muc goc, hoac discover -s tests)"""
import copy
import sys
import tempfile
import threading
import time as time_mod
import unittest
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from sqlalchemy import func, select  # noqa: E402

from app_helpers import SETTINGS, TEST_DB, TZ, add_admin, add_shift, add_worker, luot, nhan_dien, reset_db, vn  # noqa: E402
from test_app import AppCase  # noqa: E402
from test_cham_cong import ServiceCase  # noqa: E402
from app.db import Database  # noqa: E402
from app.models import Attendance, AttendanceSegment, NhatKyChamCong, PpeCheck, Worker  # noqa: E402
from app.services import bao_cao, cham_cong, cong_nhan, phep  # noqa: E402
from app.services.cham_cong import (BO_QUA_TRUNG, GHI_RA, GHI_VAO, KHONG_GHI, NGOAI_GIO_VAO, ChamCongError,  # noqa: E402
                                    danh_dau_thieu_gio_ra, sua_phien, them_phien)
from app.services.luot import DO, VANG, XANH, xu_ly_luot  # noqa: E402
from app.services import tinh_cong as tc  # noqa: E402
from src.checkpoint.evaluate import DAT, KHONG_CO_NGUOI, KHONG_DAT, NHIEU_NGUOI  # noqa: E402

ATT = SETTINGS["attendance"]
D29 = date(2026, 10, 29)  # thu Nam


class V2Case(ServiceCase):
    """ServiceCase (1 admin, ca mau, An = mat 0, Binh = mat 1) + tien ich doc dong bang cong / Viec can xu ly."""

    def ngay(self, worker=None, d=D29, now=None) -> dict:
        self.s.expire_all()
        return bao_cao.chi_tiet_ngay(self.s, self.s.get(Worker, (worker or self.an).id), d, ATT,
                                     now or datetime.combine(d, time(23, 0), tzinfo=TZ))

    def viec(self, now=None) -> list[dict]:
        self.s.expire_all()
        return bao_cao.viec_can_xu_ly(self.s, ATT, now or vn(29, 23, 0))

    def nhat_ky(self, hanh_dong=None) -> list[NhatKyChamCong]:
        q = select(NhatKyChamCong).order_by(NhatKyChamCong.id)
        if hanh_dong:
            q = q.where(NhatKyChamCong.hanh_dong == hanh_dong)
        return self.s.scalars(q).all()

    def phien(self, d=29) -> list[AttendanceSegment]:
        a = self.att(self.an, d)
        return [] if a is None else sorted(a.segments, key=lambda s: s.vao_at)


# ============================== T1-T13: vao / ra, phien, muon, ve som ==============================
class VaoRaV2Test(V2Case):
    def test_t1_vao_som_ra_dung_gio(self):
        self.run_luot(vn(29, 6, 55))
        self.run_luot(vn(29, 17, 0))
        r = self.ngay()
        self.assertEqual(len(r["phien"]), 1)
        self.assertEqual((r["di_muon_phut"], r["ve_som_phut"], r["so_phut_lam"], r["trang_thai_ngay"], r["ngay_cong"]),
                         (0, 0, 540, tc.TT_DU_CONG, 1))

    def test_t2_trong_an_han(self):
        self.run_luot(vn(29, 7, 4))
        self.run_luot(vn(29, 17, 0))
        r = self.ngay()
        self.assertEqual((r["di_muon_phut"], r["so_phut_lam"], r["trang_thai_ngay"]), (0, 536, tc.TT_DU_CONG))

    def test_t3_qua_an_han_1_giay_lam_tron_len(self):
        self.run_luot(vn(29, 7, 5, 59))
        self.assertEqual(self.att(self.an).late_minutes, 6)
        self.run_luot(vn(29, 17, 0))
        r = self.ngay()
        self.assertEqual((r["di_muon_phut"], r["trang_thai_ngay"]), (6, tc.TT_DI_MUON))

    def test_t3b_dung_bien_an_han_khong_muon(self):
        self.run_luot(vn(29, 7, 5, 0))
        self.assertEqual(self.att(self.an).late_minutes, 0)

    def test_t4_muon_6_phut(self):
        self.run_luot(vn(29, 7, 6))
        self.assertEqual(self.att(self.an).late_minutes, 6)

    def test_t5_khong_dat_roi_dat(self):
        self.run_luot(vn(29, 7, 2), KHONG_DAT, ["thieu_mu"])
        self.run_luot(vn(29, 7, 10))
        a = self.att(self.an)
        self.assertEqual((a.check_in, a.late_minutes), (vn(29, 7, 10), 10))

    def test_t6_ngoai_cua_so_vao(self):
        ev = self.run_luot(vn(29, 5, 30))
        self.assertEqual((ev["hanh_dong"], ev["ly_do_ghi"], ev["mau"]), (KHONG_GHI, NGOAI_GIO_VAO, VANG))
        self.assertEqual(ev["thong_bao"], "Chưa đến giờ vào ca (từ 06:00)")
        self.assertEqual(self.check(ev).action_reason, NGOAI_GIO_VAO)
        self.assertIsNone(self.att(self.an))
        self.assertEqual(self.run_luot(vn(29, 6, 0, 0))["hanh_dong"], GHI_VAO)  # dung 06:00:00: duoc ghi

    def test_t7_ve_som(self):
        self.run_luot(vn(29, 7, 0))
        self.run_luot(vn(29, 16, 30))
        r = self.ngay()
        self.assertEqual((r["ve_som_phut"], r["so_phut_lam"], r["trang_thai_ngay"]), (30, 510, tc.TT_VE_SOM))

    def test_t8_o_lai_muon_khong_tinh_ngoai_gio_ca(self):
        self.run_luot(vn(29, 7, 0))
        self.run_luot(vn(29, 19, 30))
        self.assertEqual(self.ngay()["so_phut_lam"], 540)

    def test_t9_ra_an_trua_dung_gio_nghi(self):
        evs = [self.run_luot(t) for t in (vn(29, 7, 0), vn(29, 11, 30), vn(29, 12, 30), vn(29, 17, 0))]
        self.assertEqual([e["vao_lai"] for e in evs], [False, False, True, False])
        r = self.ngay()
        self.assertEqual((len(r["phien"]), r["so_phut_lam"], r["phut_ra_ngoai"], r["co_ra_ngoai"]), (2, 540, 0, False))

    def test_t10_nghi_trua_qua_gio(self):
        for t in (vn(29, 7, 0), vn(29, 11, 30), vn(29, 14, 0), vn(29, 17, 0)):
            self.run_luot(t)
        r = self.ngay()
        self.assertEqual((r["so_phut_lam"], r["phut_ra_ngoai"], r["so_lan_ra_ngoai"]), (450, 90, 1))

    def test_t11_ra_ngoai_buoi_sang(self):
        for t in (vn(29, 7, 0), vn(29, 7, 30), vn(29, 8, 0), vn(29, 17, 0)):
            self.run_luot(t)
        r = self.ngay()
        self.assertEqual((r["so_phut_lam"], r["phut_ra_ngoai"]), (510, 30))

    def test_t12_chong_trung_sau_luot_vao(self):
        self.run_luot(vn(29, 7, 0))
        self.assertEqual(self.run_luot(vn(29, 7, 4))["hanh_dong"], BO_QUA_TRUNG)
        self.assertEqual(self.run_luot(vn(29, 7, 11))["hanh_dong"], GHI_RA)

    def test_t13_chong_trung_sau_luot_ra_va_het_gio_vao(self):
        self.run_luot(vn(29, 7, 0))
        self.run_luot(vn(29, 17, 0))
        ev = self.run_luot(vn(29, 17, 1))
        self.assertEqual((ev["hanh_dong"], ev["thong_bao"]),
                         (BO_QUA_TRUNG, "Nguyễn Văn An: đã ghi giờ ra 17:00, không ghi lặp"))
        ev = self.run_luot(vn(29, 17, 20))
        self.assertEqual((ev["hanh_dong"], ev["thong_bao"]), (KHONG_GHI, "Đã hết giờ vào ca hôm nay"))
        r = self.ngay()
        self.assertEqual((r["gio_ra"], r["so_phut_lam"]), ("17:00", 540))


# ============================== T14-T16: thieu gio ra, ngoai gio, khong ghi ==============================
class ThieuGioRaV2Test(V2Case):
    def test_t14_thieu_gio_ra_luu_va_xu_ly(self):
        self.run_luot(vn(29, 7, 0))
        self.assertEqual(danh_dau_thieu_gio_ra(self.s, ATT, vn(29, 21, 0, 0)), 0)
        self.assertEqual(self.ngay(now=vn(29, 21, 0))["trang_thai_ngay"], tc.TT_DANG_LAM)
        self.assertEqual(danh_dau_thieu_gio_ra(self.s, ATT, vn(29, 21, 0, 1)), 1)
        self.assertEqual(danh_dau_thieu_gio_ra(self.s, ATT, vn(29, 21, 30)), 0)  # chay lai: khong sai
        self.s.expire_all()
        self.assertEqual(self.phien()[0].trang_thai, tc.THIEU_GIO_RA)  # LUU trong CSDL
        r = self.ngay()
        self.assertEqual((r["trang_thai_ngay"], r["so_phut_lam"], r["ngay_cong"], r["thieu_gio_ra"]),
                         (tc.TT_THIEU_GIO_RA, 0, 0, True))
        self.assertEqual([(v["loai"], v["ma"], v["ngay"], v["segment_id"], v["phien_vao"]) for v in self.viec()],
                         [("thieu_gio_ra", "CN001", "2026-10-29", self.phien()[0].id, "07:00")])
        self.assertEqual(self.run_luot(vn(30, 7, 0))["hanh_dong"], GHI_VAO)  # hom sau: dong moi binh thuong
        self.assertEqual(self.att(self.an, 30).check_in, vn(30, 7, 0))
        self.assertEqual(self.phien()[0].trang_thai, tc.THIEU_GIO_RA)  # ngay cu khong bi dung
        sua_phien(self.s, self.phien()[0], None, time(17, 0), "Tổ trưởng xác nhận ra 17:00", self.admin, ATT,
                  now=vn(30, 9, 0))
        r = self.ngay(now=vn(30, 9, 0))
        self.assertEqual((r["trang_thai_ngay"], r["so_phut_lam"], r["ngay_cong"]), (tc.TT_DU_CONG, 540, 1))
        self.assertEqual([(n.hanh_dong, n.gia_tri_cu, n.gia_tri_moi) for n in self.nhat_ky()],
                         [("sua_ra", "(thiếu giờ ra)", "17:00:00")])
        self.assertEqual(self.viec(vn(30, 9, 0)), [])

    def test_t15_ngoai_gio_khong_tao_dong(self):
        for t in (vn(29, 22, 0), vn(29, 23, 50), vn(30, 0, 20)):
            with self.subTest(t=t):
                self.assertEqual(self.run_luot(t)["hanh_dong"], KHONG_GHI)
        self.assertEqual(self.s.scalar(select(func.count(Attendance.id))), 0)

    def test_t16_cac_truong_hop_khong_ghi_va_ghi_ra(self):
        self.assertEqual(self.run_luot(vn(29, 7, 0), DAT, face_index=50)["hanh_dong"], KHONG_GHI)  # nguoi la
        self.assertEqual(self.run_luot(vn(29, 7, 1), NHIEU_NGUOI)["hanh_dong"], KHONG_GHI)
        self.run_luot(vn(29, 7, 2))                               # An vao
        self.run_luot(vn(29, 7, 3), face_index=1)                 # Binh vao
        ev = self.run_luot(vn(29, 17, 0), KHONG_CO_NGUOI)         # An dang co mat + KHONG_CO_NGUOI
        self.assertEqual((ev["hanh_dong"], ev["mau"], ev["thong_bao"]),
                         (GHI_RA, XANH, "Tạm biệt Nguyễn Văn An. Đã ghi giờ ra 17:00"))
        ev = self.run_luot(vn(29, 17, 1), KHONG_DAT, ["thieu_mu"], face_index=1)  # Binh ra da thao mu
        self.assertEqual((ev["hanh_dong"], ev["mau"]), (GHI_RA, XANH))


# ============================== T17-T21: cham tay, nhat ky, nghi viec, dong y ==============================
class ChamTayV2Test(V2Case):
    def test_t17_them_phien_mo_roi_quet_ra(self):
        seg = them_phien(self.s, self.an, D29, time(7, 0), None, "Camera hỏng buổi sáng", self.admin, ATT,
                         now=vn(29, 8, 0))
        self.assertEqual((seg.trang_thai, seg.nguon_vao), (tc.MO, tc.THU_CONG))
        ev = self.run_luot(vn(29, 17, 0))
        self.assertEqual(ev["hanh_dong"], GHI_RA)
        p = self.phien()[0]
        self.assertEqual((p.trang_thai, p.ra_at, p.nguon_ra), (tc.DONG, vn(29, 17, 0), tc.TU_DONG))
        self.assertEqual([(n.hanh_dong, n.gia_tri_moi) for n in self.nhat_ky()], [("them_phien", "07:00:00–(trống)")])
        self.assertTrue(self.ngay()["co_sua_tay"])

    def test_t18_cac_truong_hop_bi_tu_choi(self):
        self.run_luot(vn(29, 7, 0))
        seg = self.phien()[0]
        cases = [
            ("giờ ra chưa tới", lambda: sua_phien(self.s, seg, None, time(17, 0), "Nhập trước", self.admin, ATT,
                                                   now=vn(29, 9, 0)), "chưa tới"),
            ("ra = vào", lambda: them_phien(self.s, self.an, date(2026, 10, 28), time(8, 0), time(8, 0), "x",
                                            self.admin, ATT, now=vn(29, 9, 0)), "Giờ ra phải sau giờ vào"),
            ("chồng lấn", lambda: them_phien(self.s, self.an, D29, time(6, 30), time(7, 30), "x", self.admin, ATT,
                                             now=vn(29, 9, 0)), "chồng lên"),
            ("lý do rỗng", lambda: them_phien(self.s, self.an, date(2026, 10, 28), time(7, 0), time(8, 0), "  ",
                                              self.admin, ATT, now=vn(29, 9, 0)), "lý do"),
            ("thêm phiên mở thứ 2", lambda: them_phien(self.s, self.an, D29, time(8, 30), None, "x", self.admin, ATT,
                                                       now=vn(29, 9, 0)), "đang mở"),
            ("ngoài thời gian làm việc", lambda: them_phien(self.s, self.an, date(2025, 12, 31), time(7, 0),
                                                            time(8, 0), "x", self.admin, ATT, now=vn(29, 9, 0)),
             "ngoài thời gian làm việc"),
        ]
        for name, fn, msg in cases:
            with self.subTest(name), self.assertRaises(ChamCongError) as cm:
                fn()
            self.assertIn(msg, str(cm.exception))
        self.assertEqual(self.nhat_ky(), [])
        self.assertEqual([(p.vao_at, p.ra_at, p.trang_thai) for p in self.phien()], [(vn(29, 7, 0), None, tc.MO)])

    def test_t19_sua_gio_vao_co_nhat_ky(self):
        self.run_luot(vn(29, 7, 0))
        sua_phien(self.s, self.phien()[0], time(6, 50), None, "Đến từ 06:50, camera nhận chậm", self.admin, ATT,
                  now=vn(29, 9, 0))
        n = self.nhat_ky()
        self.assertEqual([(x.hanh_dong, x.gia_tri_cu, x.gia_tri_moi, x.user_id, x.ly_do) for x in n],
                         [("sua_vao", "07:00:00", "06:50:00", self.admin.id, "Đến từ 06:50, camera nhận chậm")])
        self.assertEqual(self.att(self.an).check_in, vn(29, 6, 50))
        self.run_luot(vn(29, 17, 0))  # luot tu dong sau do khong ghi de thong tin sua tay
        a = self.att(self.an)
        self.assertEqual((a.updated_by, a.updated_at, a.status), (self.admin.id, vn(29, 9, 0), "cham_thu_cong"))

    def test_t20_nghi_viec_khi_dang_co_mat(self):
        self.run_luot(vn(29, 7, 0))
        cong_nhan.cho_nghi_viec(self.s, self.an, self.dk, self.admin, now=vn(29, 10, 0))
        p = self.phien()[0]
        self.assertEqual((p.ra_at, p.nguon_ra, p.trang_thai), (vn(29, 10, 0), tc.THU_CONG, tc.DONG))
        self.assertEqual(self.att(self.an).worked_minutes, 180)
        self.assertEqual(self.s.get(Worker, self.an.id).ngay_nghi_viec, D29)
        self.assertEqual([n.hanh_dong for n in self.nhat_ky()], ["dong_khi_nghi_viec"])
        self.assertEqual(self.run_luot(vn(29, 17, 0))["trang_thai_mat"], "KHONG_NHAN_RA")
        self.assertEqual(self.viec(), [])  # dong phien luc cho nghi viec: khong phai "ve som chua co phep"
        r = self.ngay(d=date(2026, 10, 30), now=vn(31, 9, 0))
        self.assertEqual(r["trang_thai_ngay"], tc.TT_NGOAI_THOI_GIAN)  # sau ngay nghi viec: khong tinh vang
        cong_nhan.di_lam_lai(self.s, self.an, self.dk, now=vn(31, 8, 0))
        w = self.s.get(Worker, self.an.id)
        self.assertEqual((w.ngay_nghi_viec, w.ngay_vao_lam, w.consent_at), (None, date(2026, 10, 31), None))

    def test_sua_ca_khong_doi_so_lieu_ngay_cu(self):
        """Gio ca duoc chep vao dong cong luc tao: sua ca chi tinh lai hom nay, ngay cu giu nguyen."""
        for t in (vn(28, 7, 10), vn(28, 17, 0), vn(29, 7, 10)):
            self.run_luot(t)
        cham_cong.sua_ca(self.s, self.shift, {"start_time": time(7, 30)}, ATT, now=vn(29, 9, 0))
        self.run_luot(vn(29, 17, 0))
        hom_qua, hom_nay = self.ngay(d=date(2026, 10, 28)), self.ngay()
        self.assertEqual((hom_qua["di_muon_phut"], hom_qua["so_phut_lam"], hom_qua["trang_thai_ngay"]),
                         (10, 530, tc.TT_DI_MUON))
        self.assertEqual((hom_nay["di_muon_phut"], hom_nay["so_phut_lam"], hom_nay["trang_thai_ngay"]),
                         (0, 510, tc.TT_DU_CONG))

    def test_t21_chua_dong_y_khong_so_khop(self):
        add_worker(self.s, "CN030", "Đặng Văn Hải", face_index=30, consent=False)  # co 3 mau, chua dong y
        self.nd.xoa_bo_nho_dem()
        self.assertEqual(self.run_luot(vn(29, 7, 0), face_index=30)["trang_thai_mat"], "KHONG_NHAN_RA")


# ============================== T22: ghi dong thoi ==============================
class _MoiLanMotKhoa(dict):
    """Thay cham_cong._KHOA: moi lan lay khoa tra ve 1 khoa MOI -> bo khoa trong tien trinh, chi con khoa CSDL."""

    def setdefault(self, key, default=None):
        return threading.RLock()


class DongThoiT22Test(unittest.TestCase):
    """2 luot cua cung 1 nguoi cach nhau 0,5 giay, 2 thread, moi thread 1 session. SQLite: file tam (CSDL trong bo nho
    dung chung 1 ket noi nen khong song song duoc); MariaDB / PostgreSQL: TEST_DATABASE_URL.
    Cua so tranh chap duoc NOI RONG (cho 0,2 s sau khi doc cac phien) de test bat duoc loi neu thieu khoa: bo khoa thi
    ngay da co dong cong (vao lai sau nghi trua) se sinh 2 phien mo."""
    LAN = 5

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        url = TEST_DB if not TEST_DB.startswith("sqlite") else "sqlite:///" + (Path(self.tmp.name) / "t22.db").as_posix()
        reset_db(url)
        self.db = Database(url)
        self.db.create_all()
        with self.db.session() as s:
            add_admin(s)
            add_shift(s)
            self.wid = add_worker(s, "CN001", "Nguyễn Văn An", face_index=0).id
        self.nd = nhan_dien()

    def tearDown(self):
        self.db.engine.dispose()
        self.tmp.cleanup()

    def mot_luot(self, t, s=None) -> str:
        with s or self.db.session() as s:
            return xu_ly_luot(s, luot(t, DAT, (), 0), settings=SETTINGS, nhan_dien=self.nd,
                              image_dir=Path(self.tmp.name) / "anh")["hanh_dong"]

    def chay(self, t0: datetime) -> tuple[list, list]:
        """2 luot t0 va t0 + 0,5 s chay song song (Barrier); cho 0,2 s sau khi doc phien de noi rong cua so tranh chap."""
        barrier, res, errs = threading.Barrier(2), [], []
        cham = cham_cong.moc_gan_nhat

        def moc_cham(segs):
            time_mod.sleep(0.2)
            return cham(segs)

        def chay_luot(t):
            with self.db.session() as s:
                barrier.wait()
                try:
                    res.append(self.mot_luot(t, s))
                except Exception as e:  # noqa: BLE001 - ghi lai de bao loi
                    errs.append(repr(e))
        with mock.patch.object(cham_cong, "moc_gan_nhat", moc_cham):
            th = [threading.Thread(target=chay_luot, args=(t,)) for t in (t0, t0 + timedelta(milliseconds=500))]
            for t in th:
                t.start()
            for t in th:
                t.join(60)
        return sorted(res), errs

    def kiem_tra(self, d: date, res: list, errs: list, so_luot: int, so_phien: int):
        self.assertEqual(errs, [])
        self.assertEqual(res, [BO_QUA_TRUNG, GHI_VAO])
        with self.db.session() as s:
            start, end = cham_cong.khoang_ngay(d, TZ)
            self.assertEqual(s.scalar(select(func.count(PpeCheck.id)).where(PpeCheck.checked_at >= start,
                                                                            PpeCheck.checked_at < end)), so_luot)
            att = s.scalars(select(Attendance).where(Attendance.work_date == d)).all()
            self.assertEqual((len(att), len(att[0].segments)), (1, so_phien))
            self.assertEqual(sum(x.trang_thai == tc.MO for x in att[0].segments), 1)

    def chay_ca_hai(self, ngay0: int):
        for i in range(self.LAN):
            with self.subTest(vao_dau_ngay=i):  # ngay chua co dong cong: 2 luot cung luc 07:00
                d = date(2026, 10, ngay0 + i)
                self.kiem_tra(d, *self.chay(vn(d.day, 7, 0)), so_luot=2, so_phien=1)
            with self.subTest(vao_lai_sau_trua=i):  # ngay DA co dong cong (phien sang da dong): 2 luot cung luc 12:30
                d = date(2026, 11, 2 + i)
                self.mot_luot(vn(d.day, 7, 0, month=11))
                self.mot_luot(vn(d.day, 11, 30, month=11))
                self.kiem_tra(d, *self.chay(vn(d.day, 12, 30, month=11)), so_luot=4, so_phien=2)

    def test_t22_hai_luot_song_song(self):
        self.chay_ca_hai(12)

    def test_t22b_chi_con_khoa_csdl(self):
        """Bo khoa trong tien trinh (nhu 2 tien trinh khac nhau): khoa dong workers (FOR UPDATE) + doc lai sau khoa."""
        with mock.patch.object(cham_cong, "_KHOA", _MoiLanMotKhoa()):
            self.chay_ca_hai(19)


# ============================== T23: thoi diem luot ==============================
class ThoiDiemLuotT23Test(unittest.TestCase):
    def test_t23_worker_ghi_luc_bat_dau_luot(self):
        from fakes import CAM_CFG, FULL, ZONE, FakeCapture
        from fakes import SETTINGS as GATE_SETTINGS
        from src.checkpoint.camera import CameraReader
        from src.gate.worker import GateWorker
        from test_face import HEAD_FACE
        from test_gate import AlwaysDetector, FakeEmbedder, FakeFaceDetector

        s = copy.deepcopy(GATE_SETTINGS)
        s["checkpoint"]["timeout_s"] = 0.3
        s["gate"].update(idle_max_fps=200, trigger_frames=3, cooldown_seconds=4)
        calls = []
        with CameraReader("0", CAM_CFG, capture_factory=lambda src: FakeCapture(read_delay=0.01)) as cam:
            worker = GateWorker(cam, ZONE, s, AlwaysDetector(FULL), FakeFaceDetector([HEAD_FACE]), FakeEmbedder(),
                                on_result=lambda r: calls.append((r, datetime.now(timezone.utc))) or True)
            t = time_mod.monotonic()
            while not calls and time_mod.monotonic() - t < 5:
                worker.step()
        r, luc_co_ket_qua = calls[0]
        self.assertLess(r["bat_dau_luc"], luc_co_ket_qua)
        self.assertGreaterEqual((luc_co_ket_qua - r["bat_dau_luc"]).total_seconds(), 0.25)  # ca luot ~0,3 s

    def test_t23_gate_service_dung_luc_bat_dau(self):
        from app.events import Broadcaster
        from app.gate_service import GateService
        from app_helpers import app_config

        captured = []

        def fake_xu_ly(s, l, **kw):
            captured.append(l.thoi_diem)
            return {"hanh_dong": KHONG_GHI}
        cfg = app_config()
        gate = GateService(cfg, Database("sqlite://"), nhan_dien(), Broadcaster(), tempfile.gettempdir())
        bat_dau = datetime(2026, 10, 29, 0, 0, 1, tzinfo=timezone.utc)
        with mock.patch("app.gate_service.xu_ly_luot", fake_xu_ly), mock.patch.object(gate.events, "publish_threadsafe"):
            gate._on_result({"ket_qua": DAT, "khuon_mat": {"vectors": []}, "bat_dau_luc": bat_dau})
        self.assertEqual(captured, [bat_dau])


# ============================== T24-T25: tong quan, bang cong thang ==============================
class BaoCaoV2Test(V2Case):
    def test_t24_luot_khong_dat_chi_tinh_man_hinh_do(self):
        self.run_luot(vn(29, 7, 0), KHONG_DAT, ["thieu_mu"])            # An vao khong dat: DO
        self.run_luot(vn(29, 7, 1), KHONG_DAT, ["thieu_mu"], face_index=50)  # nguoi la khong dat: VANG
        self.run_luot(vn(29, 7, 2), face_index=1)                       # Binh vao
        ev = self.run_luot(vn(29, 16, 0), KHONG_DAT, ["thieu_mu"], face_index=1)  # Binh ra da thao mu: XANH
        self.assertEqual(ev["mau"], XANH)
        tq = bao_cao.tong_quan(self.s, ATT, vn(29, 18, 0))
        self.assertEqual((tq["luot_khong_dat"], tq["luot_dat"], tq["ty_le_dat"]), (1, 1, 0.5))
        gio7 = tq["luot_theo_gio"][7]
        self.assertEqual((gio7[XANH], gio7[DO], gio7[VANG]), (1, 1, 1))

    def test_t25_trang_thai_tung_ngay_trong_thang(self):
        w = add_worker(self.s, "CN040", "Phan Văn Lâm", face_index=40, ngay_vao_lam=date(2026, 10, 10))
        self.nd.xoa_bo_nho_dem()
        for t in (vn(12, 7, 0), vn(12, 17, 0)):                         # thu Hai 12/10: du cong
            self.run_luot(t, face_index=40)
        th = bao_cao.bang_cong_thang(self.s, 2026, 10, ATT, vn(29, 12, 0))
        r = next(x for x in th["ds"] if x["ma"] == "CN040")
        tt = r["trang_thai_ngay"]
        self.assertEqual((tt[5], tt[11], tt[12], tt[13], tt[29], tt[30]),
                         (tc.TT_NGOAI_THOI_GIAN, tc.TT_NGHI_TUAN, tc.TT_DU_CONG, tc.TT_VANG, tc.TT_CHUA_VAO,
                          tc.TT_CHUA_DEN))
        # 10/10 (thu Bay) -> 28/10: 19 ngay, tru 3 Chu nhat (11, 18, 25) va 12/10 co cong -> 15 ngay vang
        self.assertEqual((r["so_ngay_cong"], r["so_ngay_vang"], r["tong_phut"]), (1, 15, 540))
        self.assertEqual(list(r["ngay"]), [12])  # "ngay": chi ngay co dong cong (giu nhu truoc)
        data = bao_cao.excel_bang_cong_thang(th)
        from openpyxl import load_workbook
        import io
        ws = load_workbook(io.BytesIO(data))["Bang cong"]
        header = [c.value for c in ws[2]]
        self.assertEqual(header[-13:-10], ["Số ngày công", "Tổng giờ làm", "Số lần đi muộn"])
        self.assertIn("Công quy đổi", header)


# ============================== P1-P7: phep / giai trinh ==============================
class PhepTest(V2Case):
    def phep(self, loai, d=D29, now=None, **kw):
        return phep.them_phep(self.s, self.an, d, loai, "Tổ trưởng cho phép", self.admin, ATT,
                              now=now or vn(29, 6, 0), **kw)

    def test_p1_ve_som_co_phep(self):
        self.phep(tc.VE_SOM, tu_gio=time(15, 0))
        self.run_luot(vn(29, 7, 2))
        ev = self.run_luot(vn(29, 15, 2))
        self.assertEqual(ev["thong_bao"], "Tạm biệt Nguyễn Văn An. Đã ghi giờ ra 15:02 (về sớm có phép)")
        self.assertTrue(ev["som_co_phep"])
        r = self.ngay()
        self.assertEqual((r["trang_thai_ngay"], r["ve_som_phut"], r["so_phut_lam"]), (tc.TT_VE_SOM_CO_PHEP, 118, 420))
        self.assertEqual(self.viec(), [])

    def test_p2_ve_som_qua_muc_phep(self):
        self.phep(tc.VE_SOM, tu_gio=time(15, 0))
        self.run_luot(vn(29, 7, 0))
        self.run_luot(vn(29, 13, 0))
        self.assertEqual(self.ngay()["trang_thai_ngay"], tc.TT_VE_SOM)
        self.assertEqual([v["loai"] for v in self.viec()], ["ve_som_khong_phep"])

    def test_p3_them_phep_sau(self):
        self.run_luot(vn(29, 7, 0))
        ev = self.run_luot(vn(29, 16, 30))
        self.assertEqual(ev["thong_bao_phu"], "Về sớm chưa có phép, vui lòng báo tổ trưởng")
        self.assertEqual(self.ngay()["trang_thai_ngay"], tc.TT_VE_SOM)
        self.assertEqual([v["loai"] for v in self.viec()], ["ve_som_khong_phep"])
        self.phep(tc.VE_SOM, tu_gio=time(16, 30), now=vn(29, 18, 0))
        self.assertEqual(self.ngay()["trang_thai_ngay"], tc.TT_VE_SOM_CO_PHEP)
        self.assertEqual(self.viec(), [])
        self.assertEqual([n.hanh_dong for n in self.nhat_ky()], ["them_phep"])

    def test_p4_ve_som_duoi_nguong(self):
        self.run_luot(vn(29, 7, 0))
        ev = self.run_luot(vn(29, 16, 50))
        self.assertIsNone(ev["thong_bao_phu"])
        self.assertEqual(self.ngay()["trang_thai_ngay"], tc.TT_VE_SOM)
        self.assertEqual(self.viec(), [])

    def test_p5_ra_ngoai_cong_viec_tinh_cong(self):
        self.phep(tc.RA_NGOAI_CONG_VIEC, tu_gio=time(9, 0), den_gio=time(11, 0))
        for t in (vn(29, 7, 0), vn(29, 9, 0), vn(29, 11, 0), vn(29, 17, 0)):
            ev = self.run_luot(t)
        r = self.ngay()
        self.assertEqual((r["so_phut_lam"], r["phut_cong_tac"], r["phut_ra_ngoai"], r["cong_quy_doi"]),
                         (420, 120, 0, 1.0))
        self.assertEqual(r["trang_thai_ngay"], tc.TT_DU_CONG)
        self.assertIsNone(ev["thong_bao_phu"])

    def test_p6_nghi_ca_ngay(self):
        self.phep(tc.NGHI_CA_NGAY, d=date(2026, 10, 28))
        r = self.ngay(d=date(2026, 10, 28), now=vn(29, 9, 0))
        self.assertEqual((r["trang_thai_ngay"], r["ngay_cong"]), (tc.TT_NGHI_PHEP, 0))
        th = bao_cao.bang_cong_thang(self.s, 2026, 10, ATT, vn(29, 9, 0))
        an = next(x for x in th["ds"] if x["ma"] == "CN001")
        self.assertEqual(an["so_ngay_nghi_phep"], 1)
        self.assertEqual(an["trang_thai_ngay"][28], tc.TT_NGHI_PHEP)

    def test_p7_di_muon_co_phep(self):
        self.phep(tc.DI_MUON, den_gio=time(9, 0))
        self.run_luot(vn(29, 8, 50))
        self.run_luot(vn(29, 17, 0))
        r = self.ngay()
        self.assertEqual((r["trang_thai_ngay"], r["di_muon_phut"]), (tc.TT_DI_MUON_CO_PHEP, 110))
        self.assertEqual(self.viec(), [])

    def test_xac_nhan_khong_phep_roi_khoi_danh_sach(self):
        self.run_luot(vn(29, 7, 30))  # muon 30 phut, khong phep
        self.run_luot(vn(29, 17, 0))
        v = self.viec()
        self.assertEqual([(x["loai"], x["phut"]) for x in v], [("di_muon_khong_phep", 30)])
        cham_cong.xac_nhan_khong_phep(self.s, self.att(self.an), None, self.admin, ATT, now=vn(29, 18, 0))
        self.assertEqual(self.viec(), [])
        self.assertEqual([n.hanh_dong for n in self.nhat_ky()], ["xac_nhan_khong_phep"])

    def test_huy_phep_va_kiem_tra(self):
        p = self.phep(tc.VE_SOM, tu_gio=time(15, 0))
        self.run_luot(vn(29, 7, 0))
        self.run_luot(vn(29, 15, 0))
        self.assertEqual(self.ngay()["trang_thai_ngay"], tc.TT_VE_SOM_CO_PHEP)
        phep.huy_phep(self.s, p, "Nhập nhầm", self.admin, ATT, now=vn(29, 18, 0))
        self.assertEqual(self.ngay()["trang_thai_ngay"], tc.TT_VE_SOM)
        self.assertEqual([n.hanh_dong for n in self.nhat_ky()], ["them_phep", "huy_phep"])
        for kw, msg in (({"loai": tc.VE_SOM}, "giờ được phép về"),
                        ({"loai": tc.NGHI_CA_NGAY, "d": date(2026, 12, 15)}, "tối đa 30 ngày"),
                        ({"loai": tc.RA_NGOAI_CONG_VIEC, "tu_gio": time(10, 0), "den_gio": time(9, 0)}, "đến giờ")):
            with self.subTest(kw=kw), self.assertRaises(ChamCongError) as cm:
                self.phep(**kw)
            self.assertIn(msg, str(cm.exception))

    def test_ra_nghi_trua_khong_nhac_ve_som(self):
        self.run_luot(vn(29, 7, 0))
        self.assertIsNone(self.run_luot(vn(29, 11, 30))["thong_bao_phu"])

    def test_bao_cao_thang_tach_co_phep_khong_phep(self):
        self.phep(tc.DI_MUON, d=date(2026, 10, 27), den_gio=time(9, 0))
        self.phep(tc.VE_SOM, d=date(2026, 10, 23), tu_gio=time(15, 0))
        self.phep(tc.RA_NGOAI_CONG_VIEC, d=date(2026, 10, 21), tu_gio=time(9, 0), den_gio=time(10, 0))
        for d, ts in ((28, [(7, 20), (17, 0)]), (27, [(8, 50), (17, 0)]), (26, [(7, 0), (16, 0)]),
                      (23, [(7, 0), (15, 0)]), (21, [(7, 0), (9, 0), (10, 0), (17, 0)])):
            for h, m in ts:
                self.run_luot(vn(d, h, m))
        an = next(x for x in bao_cao.bang_cong_thang(self.s, 2026, 10, ATT, vn(29, 9, 0))["ds"] if x["ma"] == "CN001")
        self.assertEqual((an["so_lan_muon_khong_phep"], an["so_lan_muon_co_phep"], an["so_lan_ve_som_khong_phep"],
                          an["so_lan_ve_som_co_phep"], an["tong_phut_cong_tac"]), (1, 1, 1, 1, 60))


# ============================== API ==============================
class ApiV2Test(AppCase):
    def setUp(self):
        super().setUp()
        self.login()

    def test_phien_nhat_ky_viec_can_xu_ly_phep_ca(self):
        r = self.client.post("/api/phien", json={"worker_id": self.an_id, "ngay": "2026-10-29", "gio_vao": "07:00",
                                                  "gio_ra": "16:00", "ly_do": "Camera hỏng"})
        self.assertEqual(r.status_code, 201, r.text)
        seg_id, att_id = r.json()["phien"]["id"], r.json()["ngay"]["attendance_id"]
        self.assertEqual(r.json()["ngay"]["trang_thai_ngay"], tc.TT_VE_SOM)
        r = self.client.patch(f"/api/phien/{seg_id}", json={"gio_ra": "17:00", "ly_do": "Sửa lại"})
        self.assertEqual((r.status_code, r.json()["ngay"]["trang_thai_ngay"]), (200, tc.TT_DU_CONG))
        self.assertEqual(self.client.patch(f"/api/phien/{seg_id}", json={"gio_ra": "17:00"}).status_code, 422)
        logs = self.client.get(f"/api/bang-cong/{att_id}/nhat-ky").json()["ds"]
        self.assertEqual([(x["hanh_dong"], x["gia_tri_cu"], x["gia_tri_moi"], x["nguoi_sua"]) for x in logs],
                         [("them_phien", None, "07:00:00–16:00:00", "Quản trị"), ("sua_ra", "16:00:00", "17:00:00",
                                                                                  "Quản trị")])
        r = self.client.post(f"/api/phien/{seg_id}/huy", json={"ly_do": "Nhập nhầm ngày"})
        self.assertEqual(r.json()["ngay"]["trang_thai_ngay"], tc.TT_VANG)
        r = self.client.post("/api/phep", json={"worker_id": self.an_id, "ngay": "2026-10-28", "loai": "nghi_ca_ngay",
                                                 "ly_do": "Ốm"})
        self.assertEqual((r.status_code, r.json()["mo_ta"]), (201, "Nghỉ cả ngày"))
        self.assertEqual(len(self.client.get("/api/phep", params={"worker_id": self.an_id}).json()["ds"]), 1)
        d = self.client.get(f"/api/bang-cong/cong-nhan/{self.an_id}/2026-10-28").json()
        self.assertEqual((d["trang_thai_ngay"], [x["hanh_dong"] for x in d["nhat_ky"]]),
                         (tc.TT_NGHI_PHEP, ["them_phep"]))
        self.assertEqual(self.client.get("/api/viec-can-xu-ly").json(), {"tong": 0, "ds": []})
        self.assertIn("viec_can_xu_ly", self.client.get("/api/tong-quan").json())
        r = self.client.patch("/api/ca-lam", json={"phut_chong_trung": 1})
        self.assertEqual(r.json()["phut_chong_trung"], 1)
        self.assertEqual(self.client.patch("/api/ca-lam", json={"phut_chong_trung": 0}).status_code, 422)
        w = self.client.get(f"/api/cong-nhan/{self.an_id}").json()
        self.assertEqual((w["ngay_vao_lam"], w["ngay_nghi_viec"]), ("2026-01-01", None))

    def test_ngay_nghi_tuan_sua_tren_trang_ca(self):
        self.assertEqual(self.client.get("/api/ca-lam").json()["ngay_nghi_tuan"], ["chu_nhat"])  # theo settings.yaml
        thu_bay = "2026-10-24"
        trang_thai = lambda: self.client.get(f"/api/bang-cong/cong-nhan/{self.an_id}/{thu_bay}").json()["trang_thai_ngay"]  # noqa: E731
        self.assertEqual(trang_thai(), tc.TT_VANG)
        r = self.client.patch("/api/ca-lam", json={"ngay_nghi_tuan": ["chu_nhat", "thu_bay"]})
        self.assertEqual(r.json()["ngay_nghi_tuan"], ["thu_bay", "chu_nhat"])
        self.assertEqual(trang_thai(), tc.TT_NGHI_TUAN)
        self.assertEqual(self.client.patch("/api/ca-lam", json={"ngay_nghi_tuan": ["thu_tam"]}).status_code, 422)
        self.assertEqual(self.client.patch("/api/ca-lam", json={"ngay_nghi_tuan": []}).json()["ngay_nghi_tuan"], [])
        self.assertEqual(self.client.get(f"/api/bang-cong/cong-nhan/{self.an_id}/2026-10-25").json()["trang_thai_ngay"],
                         tc.TT_VANG)  # Chu nhat cung la ngay lam khi khong dat ngay nghi

    def test_csdl_cu_thieu_cot_bao_loi_ro_rang(self):
        self.assertEqual(self.app.state.db.cot_con_thieu(), [])
        old = Database("sqlite://")  # CSDL tao tu phien ban cu: bang shifts chua co cot moi
        with old.engine.begin() as c:
            c.exec_driver_sql("CREATE TABLE shifts (id INTEGER PRIMARY KEY, name VARCHAR(50))")
        self.assertIn("shifts.phut_chong_trung", old.cot_con_thieu())
        old.engine.dispose()


if __name__ == "__main__":
    unittest.main()
