"""Test giao dien web (F3): cac trang HTML, file tinh (Bootstrap khong dung CDN) va cac API them cho giao dien
(ca lam, danh gia luot, bang cong thang + Excel, tong quan). CSDL SQLite trong bo nho, khong can camera/mo hinh.
Chay: .venv\\Scripts\\python.exe -m unittest discover -s tests -v"""
import io
import re
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from test_app import AppCase  # noqa: E402  (dung chung cach tao ung dung test + dang nhap)

from app_helpers import vn  # noqa: E402
from app.db import utcnow  # noqa: E402
from app.models import Attendance  # noqa: E402
from app.services.cham_cong import mui_gio  # noqa: E402
from src.checkpoint.evaluate import KHONG_DAT  # noqa: E402

PAGES = ["/cu/", "/cu/cong", "/cu/cong-nhan", "/cu/to-doi", "/cu/lich-su", "/cu/bang-cong", "/cu/bang-cong/thang", "/cu/ca-lam",
         "/cu/cau-hinh-cong"]  # giao dien cu (du phong)


class PagesTest(AppCase):
    def test_chua_dang_nhap_chuyen_ve_login(self):
        for path in PAGES + ["/cu/cong-nhan/1"]:
            with self.subTest(path=path):
                r = self.client.get(path, follow_redirects=False)
                self.assertEqual((r.status_code, r.headers["location"]), (303, f"/cu/login?next={path}"))

    def test_da_dang_nhap_moi_trang_200_tieng_viet_khong_cdn(self):
        self.login()
        for path in PAGES + [f"/cu/cong-nhan/{self.an_id}"]:
            with self.subTest(path=path):
                r = self.client.get(path)
                self.assertEqual(r.status_code, 200)
                self.assertIn("text/html", r.headers["content-type"])
                self.assertIn('lang="vi"', r.text)
                self.assertIn("/static/vendor/bootstrap-5.3.3/bootstrap.min.css", r.text)
                self.assertIsNone(re.search(r'(src|href)="https?://', r.text), "khong tai file tu mang ngoai (CDN)")
        menu = self.client.get("/cu/").text
        for label in ("Tổng quan", "Màn hình cổng", "Công nhân", "Tổ đội", "Lịch sử kiểm tra", "Bảng công", "Ca làm", "Cấu hình cổng",
                      "Đăng xuất", "Quản trị"):
            self.assertIn(label, menu)
        self.assertIn("Nguyễn Văn An", self.client.get(f"/cu/cong-nhan/{self.an_id}").text)
        self.assertEqual(self.client.get("/cu/cong-nhan/999").status_code, 404)

    def test_trang_dang_nhap_va_file_tinh_khong_can_dang_nhap(self):
        r = self.client.get("/cu/login")
        self.assertEqual(r.status_code, 200)
        self.assertIn("Đăng nhập", r.text)
        for path in ("/static/vendor/bootstrap-5.3.3/bootstrap.min.css", "/static/vendor/bootstrap-5.3.3/bootstrap.bundle.min.js",
                     "/static/js/common.js", "/static/js/cong.js", "/static/css/app.css"):
            with self.subTest(path=path):
                self.assertEqual(self.client.get(path).status_code, 200)
        r = self.client.post("/cu/login", data={"username": "admin", "password": "sai", "next": "/cu/"})
        self.assertEqual(r.status_code, 401)
        self.assertIn("Sai tên đăng nhập hoặc mật khẩu", r.text)

    def test_moi_file_js_trang_deu_ton_tai(self):
        self.login()
        for path in PAGES + [f"/cu/cong-nhan/{self.an_id}"]:
            for js in re.findall(r'src="(/static/[^"?]+)', self.client.get(path).text):
                with self.subTest(page=path, js=js):
                    self.assertEqual(self.client.get(js).status_code, 200)


class UiApiTest(AppCase):
    def setUp(self):
        super().setUp()
        self.login()

    def test_sua_ca_lam_kiem_tra_va_tinh_lai_hom_nay(self):
        s = self.client.get("/api/ca-lam").json()
        self.assertEqual((s["start_time"], s["end_time"], s["break_start"], s["break_minutes"]), ("07:00", "17:00", "11:30", 60))
        for bad, msg in (({"start_time": "17:00", "end_time": "07:00"}, "Giờ kết thúc"),
                         ({"break_start": "06:30"}, "Giờ nghỉ phải nằm trong giờ ca")):
            with self.subTest(bad=bad):
                r = self.client.patch("/api/ca-lam", json=bad)
                self.assertEqual(r.status_code, 400)
                self.assertIn(msg, r.json()["detail"])
        self.assertEqual(self.client.patch("/api/ca-lam", json={"grace_minutes": 500}).status_code, 422)
        # hom nay co 1 dong cong 08:00-11:00 -> doi gio nghi sang 10:00 thi bi tru 60 phut
        tz = mui_gio(self.app.state.settings["attendance"])
        today = utcnow().astimezone(tz).date()
        with self.app.state.db.session() as db:
            a = Attendance(worker_id=self.an_id, work_date=today, shift_id=1,
                           check_in=vn(today.day, 8, 0, month=today.month, year=today.year),
                           check_out=vn(today.day, 11, 0, month=today.month, year=today.year), worked_minutes=180)
            db.add(a)
            db.commit()
            aid = a.id
        r = self.client.patch("/api/ca-lam", json={"break_start": "10:00", "name": "Ca thi"}).json()
        self.assertEqual((r["break_start"], r["name"], r["so_dong_tinh_lai"]), ("10:00", "Ca thi", 1))
        with self.app.state.db.session() as db:
            self.assertEqual(db.get(Attendance, aid).worked_minutes, 120)

    def test_danh_gia_luot(self):
        ev = self.add_check(vn(29, 7, 0), KHONG_DAT, ["thieu_mu"])
        r = self.client.patch(f"/api/luot/{ev['id']}/danh-gia", json={"danh_gia": "bao_nham"})
        self.assertEqual(r.json()["danh_gia"], "bao_nham")
        self.assertEqual(self.client.patch(f"/api/luot/{ev['id']}/danh-gia", json={"danh_gia": "sai"}).status_code, 422)
        self.assertIsNone(self.client.patch(f"/api/luot/{ev['id']}/danh-gia", json={"danh_gia": None}).json()["danh_gia"])

    def test_bang_cong_thang_va_excel(self):
        self.add_check(vn(28, 7, 20))      # di muon 20 phut
        self.add_check(vn(28, 17, 0))
        self.add_check(vn(29, 7, 0))
        self.add_check(vn(29, 11, 0))      # buoi sang, khong tru nghi trua
        d = self.client.get("/api/bang-cong/thang", params={"thang": "2026-10"}).json()
        self.assertEqual((d["thang"], d["so_ngay"]), ("2026-10", 31))
        row = d["ds"][0]
        self.assertEqual((row["ma"], row["so_ngay_cong"], row["tong_phut"], row["so_lan_muon"]), ("CN001", 2, 520 + 240, 1))
        self.assertEqual(row["ngay"]["28"]["so_phut_lam"], 520)  # 07:20-17:00 = 580 - 60 nghi
        self.assertEqual(self.client.get("/api/bang-cong/thang", params={"thang": "2026-13"}).status_code, 422)
        r = self.client.get("/api/bang-cong/thang/excel", params={"thang": "2026-10"})
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r.headers["content-type"])
        self.assertIn("bang_cong_2026_10.xlsx", r.headers["content-disposition"])
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Bang cong", "Chi tiet"])
        ws = wb["Bang cong"]
        self.assertEqual(ws.cell(row=2, column=1).value, "Mã")
        self.assertEqual((ws.cell(row=3, column=1).value, ws.cell(row=3, column=3 + 28).value), ("CN001", "8:40"))
        self.assertEqual(wb["Chi tiet"].cell(row=2, column=1).value, "28/10/2026")

    def test_tong_quan_co_ty_le_dat_va_nguoi_dang_co_mat(self):
        now = utcnow().astimezone(mui_gio(self.app.state.settings["attendance"]))
        self.add_check(now, KHONG_DAT, ["thieu_mu"])
        self.add_check(now)  # DAT -> vao
        d = self.client.get("/api/tong-quan").json()
        self.assertEqual((d["luot_dat"], d["luot_khong_dat"], d["ty_le_dat"]), (1, 1, 0.5))
        self.assertEqual([p["ma"] for p in d["dang_co_mat_ds"]], ["CN001"])

    def test_cong_nhan_co_ten_nguoi_ghi_nhan_dong_y_va_bang_cong_gan_day(self):
        w = self.client.post("/api/cong-nhan", json={"code": "CN009", "full_name": "Hà Văn Kiên"}).json()
        r = self.client.post(f"/api/cong-nhan/{w['id']}/dong-y", json={"dong_y": True}).json()
        self.assertEqual(r["consent_by_name"], "Quản trị")
        now = utcnow().astimezone(mui_gio(self.app.state.settings["attendance"]))
        self.add_check(now)
        rows = self.client.get(f"/api/cong-nhan/{self.an_id}/bang-cong").json()["ds"]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["ngay"], now.date().isoformat())


if __name__ == "__main__":
    unittest.main()
