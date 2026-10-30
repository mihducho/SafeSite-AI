"""So lieu cho web (DE_TAI.md muc 9.1, docs/LOGIC_CHAM_CONG.md muc 14): bang cong theo ngay / thang (xuat Excel),
tong quan hom nay, Viec can xu ly. Trang thai moi ngay do MOT ham tinh (tinh_cong.tong_hop qua
cham_cong.tinh_tu_du_lieu), dung chung cho bang cong ngay, thang, Excel va API. Gio hien thi theo gio Viet Nam.
Cac truong cu (gio_vao, gio_ra, di_muon_phut, ve_som_phut, so_phut_lam, trang_thai, thieu_gio_ra...) giu nguyen ten;
truong moi chi them vao."""
import calendar
import io
from datetime import date, datetime, timedelta

from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session, selectinload

from app.models import DANG_LAM, Attendance, AttendanceSegment, Phep, PpeCheck, Shift, Worker
from app.services import tinh_cong as tc
from app.services.cham_cong import (GHI_VAO, ca_mac_dinh, gio_phut, khoang_ngay, moc_ca, mui_gio, ngay_cong,
                                    phep_cua_ngay, tinh_tu_du_lieu)
from app.services.luot import DO, VANG, XANH, phan_loai_luot
from app.services.phep import TEN_LOAI, mo_ta
from src.checkpoint.evaluate import DAT, KHONG_DAT


def _iso(t: datetime | None, tz) -> str | None:
    return t.astimezone(tz).isoformat(timespec="seconds") if t else None


def phien_json(s: AttendanceSegment, tz) -> dict:
    return {"id": s.id, "vao": gio_phut(s.vao_at, tz), "ra": gio_phut(s.ra_at, tz), "vao_luc": _iso(s.vao_at, tz),
            "ra_luc": _iso(s.ra_at, tz), "nguon_vao": s.nguon_vao, "nguon_ra": s.nguon_ra, "trang_thai": s.trang_thai,
            "vao_check_id": s.vao_check_id, "ra_check_id": s.ra_check_id, "ghi_chu": s.ghi_chu}


def phep_json(p: Phep, tz) -> dict:
    return {"id": p.id, "worker_id": p.worker_id, "ngay": p.ngay.isoformat(), "loai": p.loai,
            "ten_loai": TEN_LOAI.get(p.loai, p.loai), "tu_gio": f"{p.tu_gio:%H:%M}" if p.tu_gio else None,
            "den_gio": f"{p.den_gio:%H:%M}" if p.den_gio else None, "tinh_cong": p.tinh_cong, "ly_do": p.ly_do,
            "trang_thai": p.trang_thai, "nguoi_tao": p.nguoi_tao, "thoi_diem_tao": _iso(p.thoi_diem_tao, tz),
            "mo_ta": mo_ta(p)}


def dong_bang_cong(w: Worker, a: Attendance | None, shift: Shift, ngay: date, now: datetime, att_cfg: dict,
                   pheps: list[Phep] = ()) -> dict:
    """1 dong bang cong (ca ngay chua co dong cong: a = None). pheps: phep da duyet cua ngay do."""
    tz = mui_gio(att_cfg)
    segs = sorted(a.segments, key=lambda s: s.vao_at) if a is not None else []
    th = tinh_tu_du_lieu(w, ngay, a, segs, list(pheps), shift, att_cfg, now)
    return {
        "worker_id": w.id, "ma": w.code, "ho_ten": w.full_name, "to_doi": w.team.name if w.team else None,
        "ngay": ngay.isoformat(), "ca": shift.name,
        "gio_vao": gio_phut(th.check_in, tz), "gio_ra": gio_phut(th.check_out, tz),
        "di_muon_phut": th.late_minutes, "ve_som_phut": th.early_leave_minutes, "so_phut_lam": th.worked_minutes,
        "trang_thai": a.status if a is not None else "chua_vao",
        "thieu_gio_ra": th.thieu_gio_ra,
        "ghi_chu": a.note if a is not None else None,
        # cham cong v2
        "attendance_id": a.id if a is not None else None,
        "trang_thai_ngay": th.trang_thai_ngay, "ten_trang_thai": tc.TEN_TRANG_THAI[th.trang_thai_ngay],
        "ngay_cong": th.ngay_cong, "cong_quy_doi": th.cong_quy_doi,
        "so_lan_ra_ngoai": th.so_lan_ra_ngoai, "phut_ra_ngoai": th.phut_ra_ngoai, "phut_cong_tac": th.phut_cong_tac,
        "co_sua_tay": th.co_sua_tay, "co_ra_ngoai": th.co_ra_ngoai, "lam_ngay_nghi": th.lam_ngay_nghi,
        "muon_co_phep": th.muon_co_phep, "som_co_phep": th.som_co_phep, "can_giai_trinh": th.can_giai_trinh,
        "phien": [phien_json(s, tz) for s in segs],
        "phep": [phep_json(p, tz) for p in pheps],
    }


def _bang_cong(db: Session, tu: date, den: date, shift_id: int | None = None,
               worker_id: int | None = None) -> dict[tuple[int, date], Attendance]:
    q = select(Attendance).options(selectinload(Attendance.segments)).where(Attendance.work_date >= tu,
                                                                            Attendance.work_date <= den)
    if shift_id is not None:
        q = q.where(Attendance.shift_id == shift_id)
    if worker_id is not None:
        q = q.where(Attendance.worker_id == worker_id)
    return {(a.worker_id, a.work_date): a for a in db.scalars(q)}


def _phep(db: Session, tu: date, den: date) -> dict[tuple[int, date], list[Phep]]:
    out: dict[tuple[int, date], list[Phep]] = {}
    for p in db.scalars(select(Phep).where(Phep.ngay >= tu, Phep.ngay <= den, Phep.trang_thai == "da_duyet")
                        .order_by(Phep.id)):
        out.setdefault((p.worker_id, p.ngay), []).append(p)
    return out


def _workers_for(db: Session, worker_ids_with_rows, tu: date, team_id: int | None = None):
    """Cong nhan dang lam + cong nhan da nghi co cong trong khoang dang xem hoac nghi viec tu ngay `tu` tro di."""
    q = select(Worker).where(or_(Worker.status == DANG_LAM, Worker.id.in_(list(worker_ids_with_rows) or [-1]),
                                 Worker.ngay_nghi_viec >= tu))
    if team_id is not None:
        q = q.where(Worker.team_id == team_id)
    return db.scalars(q.order_by(Worker.code)).all()


def bang_cong_ngay(db: Session, ngay: date, att_cfg: dict, now: datetime, team_id: int | None = None) -> list[dict]:
    """1 dong / cong nhan dang lam (va cong nhan da nghi co cong / con lam viec ngay do), ke ca nguoi chua vao."""
    shift = ca_mac_dinh(db)
    rows = _bang_cong(db, ngay, ngay, shift.id)
    pheps = _phep(db, ngay, ngay)
    return [dong_bang_cong(w, rows.get((w.id, ngay)), shift, ngay, now, att_cfg, pheps.get((w.id, ngay), []))
            for w in _workers_for(db, {wid for wid, _ in rows}, ngay, team_id)]


def bang_cong_gan_day(db: Session, worker: Worker, att_cfg: dict, now: datetime, so_ngay: int = 14) -> list[dict]:
    """Cac ngay co dong cong cua 1 cong nhan trong so_ngay ngay gan nhat, moi nhat truoc."""
    shift = ca_mac_dinh(db)
    den = ngay_cong(now, mui_gio(att_cfg))
    tu = den - timedelta(days=so_ngay - 1)
    rows = _bang_cong(db, tu, den, worker_id=worker.id)
    pheps = _phep(db, tu, den)
    return [dong_bang_cong(worker, a, a.shift or shift, d, now, att_cfg, pheps.get((worker.id, d), []))
            for (_, d), a in sorted(rows.items(), key=lambda kv: kv[0][1], reverse=True)]


def chi_tiet_ngay(db: Session, worker: Worker, ngay: date, att_cfg: dict, now: datetime) -> dict:
    """Bang cong 1 ngay cua 1 cong nhan: dong bang cong (co cac phien, cac phep)."""
    shift = ca_mac_dinh(db)
    a = _bang_cong(db, ngay, ngay, shift.id, worker.id).get((worker.id, ngay))
    return dong_bang_cong(worker, a, shift, ngay, now, att_cfg, phep_cua_ngay(db, worker.id, ngay))


def bang_cong_thang(db: Session, nam: int, thang: int, att_cfg: dict, now: datetime, team_id: int | None = None) -> dict:
    """Bang cong 1 thang: moi cong nhan 1 dong. "ngay": o cua cac ngay CO dong cong (nhu truoc); "trang_thai_ngay":
    trang thai MOI ngay trong thang (vang, nghi tuan, chua den...); cac cot tong."""
    shift = ca_mac_dinh(db)
    so_ngay = calendar.monthrange(nam, thang)[1]
    tu, den = date(nam, thang, 1), date(nam, thang, so_ngay)
    rows, pheps = _bang_cong(db, tu, den, shift.id), _phep(db, tu, den)
    ds = []
    for w in _workers_for(db, {wid for wid, _ in rows}, tu, team_id):
        cells, trang_thai = {}, {}
        tong = dict.fromkeys(("so_ngay_cong", "tong_phut", "so_lan_muon", "so_lan_muon_khong_phep",
                              "so_lan_muon_co_phep", "so_lan_ve_som_khong_phep", "so_lan_ve_som_co_phep",
                              "so_ngay_nghi_phep", "so_ngay_vang", "so_ngay_thieu_gio_ra", "tong_phut_ra_ngoai",
                              "tong_phut_cong_tac"), 0)
        cong_quy_doi = 0.0
        for d in range(1, so_ngay + 1):
            ngay = date(nam, thang, d)
            a = rows.get((w.id, ngay))
            r = dong_bang_cong(w, a, shift, ngay, now, att_cfg, pheps.get((w.id, ngay), []))
            trang_thai[d] = r["trang_thai_ngay"]
            if a is not None:
                cells[d] = r
            tong["so_ngay_cong"] += r["ngay_cong"]
            tong["tong_phut"] += r["so_phut_lam"]
            tong["tong_phut_ra_ngoai"] += r["phut_ra_ngoai"]
            tong["tong_phut_cong_tac"] += r["phut_cong_tac"]
            cong_quy_doi += r["cong_quy_doi"]
            if r["di_muon_phut"] > 0:
                tong["so_lan_muon"] += 1
                tong["so_lan_muon_co_phep" if r["muon_co_phep"] else "so_lan_muon_khong_phep"] += 1
            if r["ve_som_phut"] > 0:
                tong["so_lan_ve_som_co_phep" if r["som_co_phep"] else "so_lan_ve_som_khong_phep"] += 1
            tong["so_ngay_nghi_phep"] += r["trang_thai_ngay"] == tc.TT_NGHI_PHEP
            tong["so_ngay_vang"] += r["trang_thai_ngay"] == tc.TT_VANG
            tong["so_ngay_thieu_gio_ra"] += r["trang_thai_ngay"] == tc.TT_THIEU_GIO_RA
        ds.append({"worker_id": w.id, "ma": w.code, "ho_ten": w.full_name, "to_doi": w.team.name if w.team else None,
                   "ngay": cells, "trang_thai_ngay": trang_thai, **tong, "tong_cong_quy_doi": round(cong_quy_doi, 2)})
    return {"thang": f"{nam:04d}-{thang:02d}", "so_ngay": so_ngay, "ca": shift.name, "ds": ds}


def _hm(minutes: int) -> str:
    return f"{minutes // 60}:{minutes % 60:02d}"


TONG_EXCEL = [("so_ngay_cong", "Số ngày công", None), ("tong_phut", "Tổng giờ làm", _hm),
              ("so_lan_muon", "Số lần đi muộn", None), ("so_lan_muon_khong_phep", "Muộn không phép", None),
              ("so_lan_muon_co_phep", "Muộn có phép", None), ("so_lan_ve_som_khong_phep", "Về sớm không phép", None),
              ("so_lan_ve_som_co_phep", "Về sớm có phép", None), ("so_ngay_nghi_phep", "Ngày nghỉ phép", None),
              ("so_ngay_vang", "Ngày vắng", None), ("so_ngay_thieu_gio_ra", "Ngày thiếu giờ ra", None),
              ("tong_phut_ra_ngoai", "Giờ ra ngoài", _hm), ("tong_phut_cong_tac", "Giờ công tác", _hm),
              ("tong_cong_quy_doi", "Công quy đổi", None)]


def excel_bang_cong_thang(data: dict) -> bytes:
    """File .xlsx: sheet "Bang cong" (moi o = so gio lam h:mm, cac cot tong) va sheet "Chi tiet" (moi ngay co cong
    1 dong, trang thai bang chu tieng Viet)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Bang cong"
    nam, thang = data["thang"].split("-")
    ws.append([f"BẢNG CÔNG THÁNG {thang}/{nam} - {data['ca']}"])
    ws["A1"].font = Font(bold=True, size=13)
    header = ["Mã", "Họ tên", "Tổ đội"] + [str(d) for d in range(1, data["so_ngay"] + 1)] + [t for _, t, _ in TONG_EXCEL]
    ws.append(header)
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
    late_fill = PatternFill("solid", fgColor="FCE4D6")
    for c in ws[2]:
        c.font, c.fill, c.alignment = bold, fill, Alignment(horizontal="center")
    for r in data["ds"]:
        row = [r["ma"], r["ho_ten"], r["to_doi"] or ""]
        for d in range(1, data["so_ngay"] + 1):
            cell = r["ngay"].get(d)
            row.append(_hm(cell["so_phut_lam"]) if cell and cell["gio_vao"] else "")
        row += [fmt(r[k]) if fmt else r[k] for k, _, fmt in TONG_EXCEL]
        ws.append(row)
        for d, cell in r["ngay"].items():  # to mau ngay di muon
            if cell["di_muon_phut"] > 0:
                ws.cell(row=ws.max_row, column=3 + d).fill = late_fill
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width, ws.column_dimensions["C"].width = 10, 24, 16
    for col in range(4, 4 + data["so_ngay"]):
        ws.column_dimensions[get_column_letter(col)].width = 6
    for col in range(4 + data["so_ngay"], 4 + data["so_ngay"] + len(TONG_EXCEL)):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.freeze_panes = "D3"

    ct = wb.create_sheet("Chi tiet")
    cols = ["Ngày", "Mã", "Họ tên", "Tổ đội", "Giờ vào", "Giờ ra", "Đi muộn (phút)", "Về sớm (phút)", "Số giờ làm",
            "Trạng thái", "Số phiên", "Ra ngoài (phút)", "Công tác (phút)", "Công quy đổi", "Ghi chú"]
    ct.append(cols)
    for c in ct[1]:
        c.font, c.fill = bold, fill
    for r in data["ds"]:
        for d, cell in sorted(r["ngay"].items()):
            y, m, dd = cell["ngay"].split("-")
            ct.append([f"{dd}/{m}/{y}", r["ma"], r["ho_ten"], r["to_doi"] or "", cell["gio_vao"] or "",
                       cell["gio_ra"] or "", cell["di_muon_phut"], cell["ve_som_phut"], _hm(cell["so_phut_lam"]),
                       cell["ten_trang_thai"], sum(p["trang_thai"] != tc.HUY for p in cell["phien"]),
                       cell["phut_ra_ngoai"], cell["phut_cong_tac"], cell["cong_quy_doi"], cell["ghi_chu"] or ""])
    for col, width in zip("ABCDEFGHIJKLMNO", (12, 10, 24, 16, 9, 9, 15, 14, 11, 22, 9, 15, 15, 13, 50)):
        ct.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- Viec can xu ly (Q9, Q14, Q21) ---
TEN_VIEC = {"thieu_gio_ra": "Thiếu giờ ra", "di_muon_khong_phep": "Đi muộn chưa có phép",
            "ve_som_khong_phep": "Về sớm chưa có phép"}


def viec_can_xu_ly(db: Session, att_cfg: dict, now: datetime) -> list[dict]:
    """Cac ngay can can bo xu ly: co phien thieu gio ra (chua tinh ngay cong); di muon / ve som KHONG phep qua
    attendance.nguong_can_giai_trinh_phut chua "xac nhan khong phep" (ve som chi xet khi ca ngay do da ket thuc)."""
    tz, nguong = mui_gio(att_cfg), att_cfg["nguong_can_giai_trinh_phut"]
    out = []

    def item(a: Attendance, loai: str, mo_ta_viec: str, phut: int | None = None) -> dict:
        return {"loai": loai, "ten_loai": TEN_VIEC[loai], "attendance_id": a.id, "worker_id": a.worker_id,
                "ma": a.worker.code, "ho_ten": a.worker.full_name, "ngay": a.work_date.isoformat(), "mo_ta": mo_ta_viec,
                "phut": phut}

    thieu = db.scalars(select(Attendance).options(selectinload(Attendance.segments)).where(
        Attendance.id.in_(select(AttendanceSegment.attendance_id).where(AttendanceSegment.trang_thai == tc.THIEU_GIO_RA))
    )).all()
    for a in thieu:
        vao = [gio_phut(s.vao_at, tz) for s in a.segments if s.trang_thai == tc.THIEU_GIO_RA]
        out.append(item(a, "thieu_gio_ra", f"Phiên vào {', '.join(vao)} chưa có giờ ra"))
    nghi_van = db.scalars(select(Attendance).options(selectinload(Attendance.segments)).where(
        or_(Attendance.late_minutes > nguong, Attendance.early_leave_minutes > nguong),
        Attendance.da_xac_nhan_khong_phep.is_(False))).all()
    for a in nghi_van:
        th = tinh_tu_du_lieu(a.worker, a.work_date, a, list(a.segments), phep_cua_ngay(db, a.worker_id, a.work_date),
                             a.shift, att_cfg, now)
        if not th.can_giai_trinh:
            continue
        if th.late_minutes > nguong and not th.muon_co_phep:
            out.append(item(a, "di_muon_khong_phep", f"Đi muộn {th.late_minutes} phút, chưa có phép", th.late_minutes))
        if th.early_leave_minutes > nguong and not th.som_co_phep and now >= moc_ca(a.shift, a.work_date, tz)[1]:
            out.append(item(a, "ve_som_khong_phep", f"Về sớm {th.early_leave_minutes} phút, chưa có phép",
                            th.early_leave_minutes))
    return sorted(out, key=lambda x: (x["ngay"], x["ma"], x["loai"]), reverse=True)


def tong_quan(db: Session, att_cfg: dict, now: datetime, latest: int = 10) -> dict:
    """So cong nhan, da vao, dang co mat (kem danh sach), luot dat / khong dat, ti le dat, so luot theo gio, Viec can
    xu ly, luot moi nhat. Mau moi luot: phan_loai_luot (giong man hinh cong)."""
    tz = mui_gio(att_cfg)
    today = ngay_cong(now, tz)
    start, end = khoang_ngay(today, tz)
    came = db.scalars(select(Attendance).options(selectinload(Attendance.segments))
                      .where(Attendance.work_date == today, Attendance.check_in.is_not(None))).all()
    present = sorted((a for a in came if any(s.trang_thai == tc.MO for s in a.segments)), key=lambda a: a.check_in)
    theo_gio = [{"gio": h, "dat": 0, "khong_dat": 0, "khac": 0, XANH: 0, DO: 0, VANG: 0} for h in range(24)]
    so_luot = luot_dat = luot_khong_dat = 0
    for t, kq, wid, action, reason in db.execute(
            select(PpeCheck.checked_at, PpeCheck.result, PpeCheck.worker_id, PpeCheck.action, PpeCheck.action_reason)
            .where(PpeCheck.checked_at >= start, PpeCheck.checked_at < end)).all():
        mau = phan_loai_luot(kq, wid is not None, action, reason)
        o = theo_gio[t.astimezone(tz).hour]
        o["dat" if kq == DAT else "khong_dat" if kq == KHONG_DAT else "khac"] += 1  # theo ket qua bao ho (nhu truoc)
        o[mau] += 1
        so_luot += 1
        luot_dat += action == GHI_VAO
        luot_khong_dat += mau == DO
    return {
        "ngay": today.isoformat(),
        "so_cong_nhan": db.scalar(select(func.count(Worker.id)).where(Worker.status == DANG_LAM)) or 0,
        "da_vao": len({a.worker_id for a in came}),
        "dang_co_mat": len({a.worker_id for a in present}),
        "so_luot": so_luot,
        # luot VAO da ghi gio vao / luot VAO cua nguoi da nhan ra bi KHONG_DAT (mau do tren man hinh cong)
        "luot_dat": luot_dat,
        "luot_khong_dat": luot_khong_dat,
        "ty_le_dat": round(luot_dat / (luot_dat + luot_khong_dat), 3) if luot_dat + luot_khong_dat else None,
        "viec_can_xu_ly": len(viec_can_xu_ly(db, att_cfg, now)),
        "dang_co_mat_ds": [{"worker_id": a.worker_id, "ma": a.worker.code, "ho_ten": a.worker.full_name,
                            "to_doi": a.worker.team.name if a.worker.team else None,
                            "gio_vao": gio_phut(a.check_in, tz)} for a in present],
        "luot_theo_gio": theo_gio,  # 24 gio: dat / khong_dat / khac (ket qua bao ho) + xanh / do / vang (mau man hinh)
        "luot_moi_nhat": db.scalars(select(PpeCheck).order_by(PpeCheck.checked_at.desc(), PpeCheck.id.desc())
                                    .limit(latest)).all(),  # API doi sang JSON
    }
